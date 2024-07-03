import os
import openpyxl
from openpyxl.utils import get_column_letter
from PyQt5.QtWidgets import QApplication, QFileDialog, QWidget, QMessageBox, QPushButton, QVBoxLayout

class Get_QuoteForm(QWidget):
    def __init__(self, name='Get_QuoteForm'):
        super(Get_QuoteForm, self).__init__()
        self.setWindowTitle(name)
        self.resize(400, 150)   # 設定彈出視窗的大小

        # 按鈕1
        self.btn_chooseFile1 = QPushButton(self)  
        self.btn_chooseFile1.setObjectName("btn_chooseFile")  
        self.btn_chooseFile1.setText("[---Choose Quote file's Path---]")
        
        # 設定視窗的佈局
        layout = QVBoxLayout()
        layout.addWidget(self.btn_chooseFile1)
        self.setLayout(layout)

        # 設定按鈕的信號與槽函數連接
        self.btn_chooseFile1.clicked.connect(self.slot_btn_chooseFile1)
        
        self.file_paths = []

    def slot_btn_chooseFile1(self):
        global QuoteName_choose
        self.file_paths, _ = QFileDialog.getOpenFileNames(self,  
                                    "Choose Quote file's Path",  
                                    "", # start path
                                    "Excel File (*.xlsx *.xls *.xlsb);;All Files (*)")   

        if self.file_paths:
            self.close()

def process_quote_file(file_path):
    # 打開 SKU_AV Summary 工作表
    wb = openpyxl.load_workbook(file_path)
    sheet = wb["SKU_AV Summary"]
    
    # 獲取最後一行
    lrow = sheet.max_row
    
    # 設置 Cat 欄位
    sheet["K1"].value = "Cat"
    formula = '=IF(IFERROR(FIND("AV#", E2, 1), -1) > 0, "AV", ' \
              'IF(IFERROR(FIND("-", E2, 1), -1) > 0, "63PN", "SKU"))'
    sheet["K2"].value = formula
    
    # 複製公式到所有行
    for row in range(2, lrow + 1):
        sheet[f"K{row}"].value = formula
    
    # 新增 Summary 工作表
    wb.create_sheet("SKU Summary", 1)
    wb.create_sheet("AV Summary", 2)
    wb.create_sheet("63PN", 3)
    wb.create_sheet("ThirdParty", 4)
    
    # 複製到 SKU Summary
    for col in range(1, 7):
        for row in range(1, lrow + 1):
            wb["SKU Summary"].cell(row=row, column=col).value = sheet.cell(row=row, column=col).value
    
    # 添加日期到 SKU Summary
    lrow2 = wb["SKU Summary"].max_row
    wb["SKU Summary"].insert_cols(2)
    wb["SKU Summary"]["B1"].value = "Date"
    wb["SKU Summary"]["B2"].value = "=TODAY()"
    for row in range(2, lrow2 + 1):
        wb["SKU Summary"].cell(row=row, column=2).number_format = "m/d/yyyy"
    
    # 複製到 AV Summary
    for row in range(1, lrow + 1):
        if sheet.cell(row=row, column=11).value == "AV":
            for col in range(1, 7):
                wb["AV Summary"].cell(row=row, column=col).value = sheet.cell(row=row, column=col).value
    
    # 複製到 63PN Summary
    for row in range(1, lrow + 1):
        if sheet.cell(row=row, column=11).value == "63PN":
            for col in range(1, 7):
                wb["63PN"].cell(row=row, column=col).value = sheet.cell(row=row, column=col).value
    
    # 複製到 ThirdParty
    wb["63PN"].insert_cols(2)
    lrow3 = wb["63PN"].max_row
    wb["ThirdParty"]["A:A"].value = wb["63PN"]["C:C"].value
    wb["ThirdParty"]["B:B"].value = wb["63PN"]["B:B"].value
    wb["ThirdParty"]["C:C"].value = wb["63PN"]["E:E"].value
    wb["ThirdParty"]["B1"].value = "Site"
    wb["ThirdParty"]["B2"].value = "'0874"
    wb["ThirdParty"]["E1"].value = "Vendor Name"
    wb["ThirdParty"]["E2"].value = "Inventec"
    
    # 保存文件
    output_path = os.path.splitext(file_path)[0] + "_rev.xlsx"
    wb.save(output_path)
    
    # 關閉舊檔案
    wb.close()
    
    return output_path

def create_docking_av_sku():
    app = QApplication([])  # 初始化 PyQt5 應用程式
    quote_form = Get_QuoteForm()  # 創建取得報價表單的視窗
    quote_form.show()  # 顯示報價表單視窗
    app.exec_()  # 開始 PyQt5 的事件迴圈，等待使用者操作完畢後返回控制

    if quote_form.file_paths:
        for file_path in quote_form.file_paths:
            try:
                output_path = process_quote_file(file_path)
                show_message_box(f"Updated file saved to:\n{output_path}\n\nDone")
            except Exception as e:
                show_message_box(f"Error processing file {file_path}:\n{str(e)}")

def show_message_box(message):
    msg_box = QMessageBox()
    msg_box.setWindowTitle("Operation Complete")
    msg_box.setText(message)
    msg_box.exec_()

if __name__ == "__main__":
    create_docking_av_sku()
