import os
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter
from PyQt5.QtWidgets import QApplication, QMessageBox, QWidget, QPushButton, QVBoxLayout, QFileDialog
from datetime import datetime

class Get_QuoteForm(QWidget):
    def __init__(self, name='Get_QuoteForm'):
        super(Get_QuoteForm, self).__init__()
        self.setWindowTitle(name)
        self.resize(400, 150)   # 設定彈出視窗的大小

        # 按鈕1
        self.btn_chooseFile1 = QPushButton(self)  
        self.btn_chooseFile1.setObjectName("btn_chooseFile")  
        self.btn_chooseFile1.setText("[---Choose Quote file's Path---]")
        
        # 設定視窗的佈局
        layout = QVBoxLayout()
        layout.addWidget(self.btn_chooseFile1)
        self.setLayout(layout)

        # 設定按鈕的信號與槽函數連接
        self.btn_chooseFile1.clicked.connect(self.slot_btn_chooseFile1)
        
        self.file_paths = []

    def slot_btn_chooseFile1(self):
        global QuoteName_choose
        start_path = os.path.join(os.getcwd(), "5. AMO")
        self.file_paths, _ = QFileDialog.getOpenFileNames(self,  
                                    "Choose Quote file's Path",  
                                    start_path, # start path
                                    "Excel File (*.xlsx *.xls *.xlsb);;All Files (*)")   

        if self.file_paths:
            self.close()

def find_summary_sheet(wb):
    """Find the first sheet containing 'summary' in its name, case insensitive."""
    for sheet_name in wb.sheetnames:
        if "summary" in sheet_name.lower():
            return wb[sheet_name]
    raise ValueError("No sheet containing 'summary' found in the workbook")

def process_quote_file(file_path):
    # Load the workbook from the provided file path
    wb = load_workbook(filename=file_path)
    
    # Check if a summary file already exists
    today_date = datetime.today().strftime('%Y-%m-%d')
    output_filename = f"Summary_AMO_{today_date}.xlsx"
    output_path = os.path.join(os.path.dirname(file_path), output_filename)
    
    # Find the summary sheet
    sheet = find_summary_sheet(wb)
    
    if os.path.exists(output_path):
        # If summary file already exists, load it instead of creating a new one
        new_wb = load_workbook(filename=output_path)
        new_sheet = new_wb.active
        
        #for row in sheet.iter_rows(values_only=True):
        #    new_sheet.append(row)
            
        # Iterate through all rows from the second row onwards in the summary sheet and copy to new workbook
        for row in sheet.iter_rows(min_row=2, values_only=True):
            new_sheet.append(row)
    else:
        # Otherwise, create a new workbook
        new_wb = Workbook()
        new_sheet = new_wb.active
        
        for row in sheet.iter_rows(values_only=True):
            new_sheet.append(row)
    
    # Save the new workbook
    new_wb.save(output_path)
    
    # Return the path where the new file is saved
    return output_path

def create_docking_av_sku():
    app = QApplication([])  # Initialize PyQt5 application
    quote_form = Get_QuoteForm()  # Create quote form window
    quote_form.show()  # Show quote form window
    app.exec_()  # Start PyQt5 event loop and wait for user interaction

    if quote_form.file_paths:
        for file_path in quote_form.file_paths:
            try:
                output_path = process_quote_file(file_path)
                show_message_box(f"Updated file saved to:\n{output_path}\n\nDone")
            except Exception as e:
                show_message_box(f"Error processing file {file_path}:\n{str(e)}")
    return output_path

def show_message_box(message):
    msg_box = QMessageBox()
    msg_box.setWindowTitle("Operation Complete")
    msg_box.setText(message)
    msg_box.exec_()

if __name__ == "__main__":
    print("hello world")
    output_path = create_docking_av_sku()
    print("created " + output_path)
